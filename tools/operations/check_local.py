"""Check local imports, GUI runtime, and Office rendering without API calls."""
from pathlib import Path
import io
import sys
import tempfile

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))


def main():
    import tkinter
    import customtkinter
    from backend.engines.artifact_pipeline import process_file
    from backend.core.model_client import create_model
    from backend.quality.validation import render_office
    import fitz
    from pptx import Presentation
    from pptx.util import Inches
    from openpyxl import Workbook
    from docx import Document

    window = tkinter.Tk()
    window.withdraw()
    window.update_idletasks()
    window.destroy()
    print('Backend, Gemini client, and desktop GUI imports: OK')
    documents = {}
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = 'Local rendering check'
    documents['pptx'] = deck
    workbook = Workbook()
    workbook.active['A1'] = 'Local rendering check'
    documents['xlsx'] = workbook
    document = Document()
    document.add_paragraph('Local rendering check')
    documents['docx'] = document
    with tempfile.TemporaryDirectory(prefix='translator-local-check-') as directory:
        for kind, document in documents.items():
            buffer = io.BytesIO()
            document.save(buffer)
            rendered = render_office(buffer.getvalue(), kind, Path(directory), kind)
            with fitz.open(stream=rendered, filetype='pdf') as pdf:
                if not any('Local rendering check' in page.get_text() for page in pdf):
                    raise RuntimeError(f'{kind} rendered without expected text')
            print(f'{kind.upper()} -> PDF rendering: OK')
    print('Local runtime checks passed. No cloud calls or document translations were made.')


if __name__ == '__main__':
    main()
